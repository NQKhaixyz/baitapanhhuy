import json, os, csv
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

here = os.path.dirname(__file__)
qa = [json.loads(l) for l in open(os.path.join(here, "..", "qa.sample.jsonl"), encoding="utf-8")]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "golden_qa"

headers = ["id", "type_id", "type", "specialty", "question",
           "retrieval_chunks", "retrieval_points", "answer_gt", "is_answerable"]
ws.append(headers)

for it in qa:
    chunks = "\n".join(cid for c in it["retrieval_gt"] for cid in c["chunk_ids"])
    pts = "\n".join("• " + p["text"] for c in it["retrieval_gt"] for p in c["points"])
    ws.append([it["id"], it["type_id"], it["type"], it["specialty"], it["question"],
               chunks, pts, it["answer_gt"], it["meta"]["is_answerable"]])

# --- styling ---
header_fill = PatternFill("solid", fgColor="0E7C86")
header_font = Font(bold=True, color="FFFFFF", size=11)
thin = Side(style="thin", color="D0DADE")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

widths = {"A": 13, "B": 8, "C": 22, "D": 12, "E": 44, "F": 20, "G": 46, "H": 52, "I": 13}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(vertical="center", horizontal="left")
ws.row_dimensions[1].height = 24
ws.freeze_panes = "A2"

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = border

wb.save(os.path.join(here, "A_flat.xlsx"))
print("Da tao A_flat.xlsx")

# --- kèm bản CSV có BOM để so sánh (Excel Mac đọc đúng dấu hơn) ---
with open(os.path.join(here, "A_flat_BOM.csv"), "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(headers)
    for it in qa:
        chunks = "|".join(cid for c in it["retrieval_gt"] for cid in c["chunk_ids"])
        pts = " ;; ".join(p["text"] for c in it["retrieval_gt"] for p in c["points"])
        w.writerow([it["id"], it["type_id"], it["type"], it["specialty"], it["question"],
                    chunks, pts, it["answer_gt"], it["meta"]["is_answerable"]])
print("Da tao A_flat_BOM.csv")
